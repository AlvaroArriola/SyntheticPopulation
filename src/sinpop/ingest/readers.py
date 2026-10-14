"""Lectores de microdatos guiados por el diseño de registro oficial (F1/F2).

Los valores se leen siempre como texto: la recodificación a variables armonizadas
se hace en F2 con las tablas de `configs/variables.yaml`.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
import polars as pl


@dataclass(frozen=True)
class Campo:
    nombre: str
    longitud: int
    diccionario: str | None
    descripcion: str


def layout_from_xlsx(path: Path, hoja: str = "Diseño") -> list[Campo]:
    """Lee la hoja de diseño (columnas Variable, Diccionario, Longitud, …) en orden."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[hoja].iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "Variable")
    hdr = [str(c).strip() if c else "" for c in rows[hdr_i]]
    i_long = hdr.index("Longitud")
    i_dic = hdr.index("Diccionario de la variable")
    i_desc = next(i for i, h in enumerate(hdr) if h.startswith("Descripci"))
    campos = []
    for r in rows[hdr_i + 1 :]:
        if not r or not r[0] or str(r[0]).startswith("***"):
            if r and r[0] and str(r[0]).startswith("***"):
                break
            continue
        campos.append(
            Campo(
                nombre=str(r[0]).strip(),
                longitud=int(r[i_long]),
                diccionario=(str(r[i_dic]).strip() if r[i_dic] else None),
                descripcion=str(r[i_desc] or "").strip(),
            )
        )
    return campos


def codelists_from_xlsx(path: Path) -> dict[str, dict[str, str]]:
    """Lee los diccionarios (hojas 'Tablas*'): bloques que empiezan con el nombre T_xxx."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict[str, str]] = {}
    for hoja in wb.sheetnames:
        if not hoja.lower().startswith("tablas"):
            continue
        actual: str | None = None
        for r in wb[hoja].iter_rows(values_only=True):
            vals = [c for c in r if c is not None and str(c).strip() != ""]
            if not vals:
                continue
            first = str(vals[0]).strip()
            if len(vals) == 1 and first.startswith("T_"):
                actual = first
                out[actual] = {}
            elif actual and first.lower().startswith("código"):
                continue
            elif actual and len(vals) >= 2:
                out[actual][first] = str(vals[1]).strip()
    return out


def read_fixed_width(
    path: Path, campos: list[Campo], columns: list[str] | None = None
) -> pl.DataFrame:
    """Lee un fichero de ancho fijo. Las líneas recortadas se rellenan con espacios."""
    total = sum(c.longitud for c in campos)
    raw = pl.read_csv(
        path,
        has_header=False,
        new_columns=["_l"],
        separator="\x1f",
        quote_char=None,
        infer_schema=False,
        encoding="utf8-lossy",
    )
    line = pl.col("_l").str.pad_end(total)
    exprs = []
    pos = 0
    for c in campos:
        if columns is None or c.nombre in columns:
            exprs.append(line.str.slice(pos, c.longitud).str.strip_chars().alias(c.nombre))
        pos += c.longitud
    return raw.select(exprs)
